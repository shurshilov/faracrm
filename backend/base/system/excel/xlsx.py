"""
Excel (.xlsx) ↔ значения полей модели.

Чистые функции без запросов к БД: сборка файла на экспорт, разбор файла на
импорт и преобразование значения поля в ячейку и обратно по типу поля.
Формат технический: заголовки — имена полей, связи и вложения — id,
варианты Selection — значения. Что без БД не решить — существование
связанных записей, связи по имени, вложения по URL — делает роутер
(routers/excel.py).
"""

import datetime
import io
import json
import re
from decimal import Decimal as PythonDecimal, InvalidOperation
from typing import Any, Iterable, Type

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from backend.base.system.dotorm.dotorm.fields import (
    BigInteger,
    Binary,
    Boolean,
    Char,
    Date,
    Datetime,
    Decimal,
    Field,
    Float,
    Integer,
    Many2many,
    Many2one,
    PolymorphicMany2one,
    Selection,
    SmallInteger,
    Text,
)
from backend.base.system.dotorm.dotorm.model import DotModel

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

# Что импорт умеет заполнить из ячейки. One2many (дети импортируются своей
# моделью с колонкой родителя), полиморфные списки, JSON и байты — нет.
IMPORTABLE_TYPES = (
    Char,  # и Selection
    Text,
    Integer,
    BigInteger,
    SmallInteger,
    Float,
    Decimal,
    Boolean,
    Date,
    Datetime,
    Many2one,
    Many2many,
    PolymorphicMany2one,
)

TRUE_WORDS = {"1", "true", "yes", "y", "да", "истина", "+"}
FALSE_WORDS = {"0", "false", "no", "n", "нет", "ложь", "-"}
DATE_FORMATS = ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y")

# Строки данных: (номер строки на листе, ячейки). Первая строка — заголовки.
Rows = list[tuple[int, list[Any]]]


def exportable_fields(Model: Type[DotModel]) -> dict[str, Field]:
    """Поля, которые можно выгрузить: публичные, кроме байтов."""
    return {
        name: field
        for name, field in Model.get_public_fields().items()
        if not isinstance(field, Binary)
    }


def importable_fields(Model: Type[DotModel]) -> dict[str, Field]:
    """Поля, которые можно заполнить из файла: не ключ, не вычисляемые,
    поддерживаемого типа (Many2many — не store, но связывается по id)."""
    return {
        name: field
        for name, field in Model.get_public_fields().items()
        if not field.primary_key
        and not field.compute
        and isinstance(field, IMPORTABLE_TYPES)
    }


# ==================== Заголовки ====================

# «uom_id (name)» — имя поля и, в скобках, поле связанной модели, по
# которому искать запись; без скобок — id.
HEADER_RE = re.compile(
    r"^\s*(.+?)\s*(?:\(\s*([A-Za-z_][A-Za-z0-9_]*)\s*\))?\s*$"
)


def parse_header(header: str) -> tuple[str, str]:
    """Заголовок колонки → (имя поля, поле поиска связи): «uom_id (name)»
    → («uom_id», «name»), «uom_id» → («uom_id», «id»)."""
    match = HEADER_RE.match(header)
    if not match:
        return header.strip(), "id"
    return match.group(1), (match.group(2) or "id").lower()


# ==================== Экспорт ====================


def _record_id(value: Any) -> Any:
    return value.id if isinstance(value, DotModel) else value


def to_cell(field: Field, value: Any) -> Any:
    """Значение поля записи → значение ячейки."""
    if value is None:
        return None
    if field.relation:
        # Связи — id: файл переносит данные и читается импортом обратно.
        if isinstance(value, list):
            return ", ".join(str(_record_id(record)) for record in value)
        return _record_id(value)
    if isinstance(value, datetime.datetime):
        # Excel не знает часовых поясов: локальное время сервера без
        # tzinfo. Импорт читает наивное время как локальное — симметрично.
        if value.tzinfo is not None:
            value = value.astimezone()
        return value.replace(tzinfo=None)
    if isinstance(
        value,
        (bool, int, float, PythonDecimal, datetime.date, datetime.time),
    ):
        return value
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    # Управляющие символы в тексте openpyxl не принимает.
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def build_workbook(headers: list[str], rows: Iterable[list[Any]]) -> bytes:
    """Один лист: жирный заголовок + строки → байты .xlsx. Write-only
    режим пишет строки потоком, не собирая лист в памяти."""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    bold = Font(bold=True)
    header_cells = []
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(header) + 4, 12), 60
        )
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = bold
        header_cells.append(cell)
    sheet.append(header_cells)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ==================== Импорт ====================


def text(value: Any) -> str:
    """Ячейка как текст: целые числа без «.0» (телефоны, коды)."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.isoformat(sep=" ")
    return str(value)


def read_sheet(content: bytes) -> tuple[list[str], Rows]:
    """Первый лист файла → заголовки (первая строка, текстом) и строки
    данных с их номерами на листе, без полностью пустых. Формулы —
    значениями (data_only). Хвостовые пустые ячейки строки openpyxl
    отрезает — читать по индексу с проверкой длины."""
    workbook = load_workbook(
        io.BytesIO(content), read_only=True, data_only=True
    )
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        headers = [
            text(cell).strip() if cell is not None else ""
            for cell in (next(rows, None) or ())
        ]
        data = [
            (row_number, list(row))
            for row_number, row in enumerate(rows, start=2)
            if any(cell is not None and str(cell).strip() for cell in row)
        ]
    finally:
        workbook.close()
    return headers, data


def read_images(content: bytes) -> dict[tuple[int, int], tuple[bytes, str]]:
    """Картинки, вставленные на первый лист: {(номер строки, индекс
    колонки) ячейки верхнего левого угла: (байты, формат)}. Рисунки
    читаются только в обычном режиме (не read_only), публичного доступа к
    ним у openpyxl нет — ws._images."""
    workbook = load_workbook(io.BytesIO(content))
    try:
        images: dict[tuple[int, int], tuple[bytes, str]] = {}
        for image in workbook.worksheets[0]._images:
            marker = getattr(image.anchor, "_from", None)
            if marker is not None:
                images.setdefault(
                    (marker.row + 1, marker.col), (image._data(), image.format)
                )
        return images
    finally:
        workbook.close()


def from_cell(field: Field, value: Any) -> Any:
    """Значение ячейки → значение поля; ValueError — не разобрать.

    Пустая ячейка → None (поле не заполняется, сработает default).
    Связи → сырые значения: число — id, текст — имя (Many2many —
    список через запятую); вложение — id или URL. В записи их
    превращает роутер.
    """
    if isinstance(value, str):
        value = value.strip()
    if value is None or value == "":
        return None
    if isinstance(field, Boolean):
        return _to_bool(value)
    if isinstance(field, (Integer, BigInteger, SmallInteger)):
        return _to_int(value)
    if isinstance(field, Float):
        return _to_float(value)
    if isinstance(field, Decimal):
        return _to_decimal(value)
    if isinstance(field, Datetime):
        return _to_datetime(value)
    if isinstance(field, Date):
        return _to_date(value)
    if isinstance(field, Selection):
        return _to_selection(field, value)
    if isinstance(field, PolymorphicMany2one):
        token = _relation_token(value)
        if isinstance(token, int) or token.lower().startswith(
            ("http://", "https://")
        ):
            return token
        raise ValueError("ожидается id вложения или URL")
    if isinstance(field, Many2many):
        tokens = [
            _relation_token(token)
            for token in text(value).split(",")
            if token.strip()
        ]
        return tokens or None
    if isinstance(field, Many2one):
        return _relation_token(value)
    return text(value)


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _relation_token(value: Any) -> int | str:
    """Ссылка на связанную запись из ячейки: число или цифры — id,
    остальное — текст (имя, URL)."""
    if _is_number(value):
        return _to_int(value)
    token = text(value).strip()
    return int(token) if token.isdigit() else token


def _to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    word = str(value).lower()
    if word in TRUE_WORDS:
        return True
    if word in FALSE_WORDS:
        return False
    raise ValueError("ожидается да/нет")


def _to_int(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError("ожидается число")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise ValueError("ожидается целое число")
    try:
        return int(str(value).replace(" ", ""))
    except ValueError:
        raise ValueError("ожидается число") from None


def _to_float(value: Any) -> float:
    if isinstance(value, bool):
        raise ValueError("ожидается число")
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except ValueError:
        raise ValueError("ожидается число") from None


def _to_decimal(value: Any) -> PythonDecimal:
    if isinstance(value, bool):
        raise ValueError("ожидается число")
    try:
        return PythonDecimal(str(value).replace(" ", "").replace(",", "."))
    except InvalidOperation:
        raise ValueError("ожидается число") from None


def _parse_datetime(value: Any) -> datetime.datetime:
    """Ячейка-дата (openpyxl отдаёт datetime) или текст: ISO либо
    ДД.ММ.ГГГГ [ЧЧ:ММ[:СС]]."""
    if isinstance(value, datetime.datetime):
        return value
    if isinstance(value, datetime.date):
        return datetime.datetime.combine(value, datetime.time())
    string = str(value)
    try:
        return datetime.datetime.fromisoformat(string)
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(string, fmt)
        except ValueError:
            continue
    raise ValueError("ожидается дата (ДД.ММ.ГГГГ)")


def _to_datetime(value: Any) -> datetime.datetime:
    parsed = _parse_datetime(value)
    # Без смещения — локальное время сервера (как у фильтров,
    # см. Datetime.to_sql_filter).
    return parsed if parsed.tzinfo else parsed.astimezone()


def _to_date(value: Any) -> datetime.date:
    return _parse_datetime(value).date()


def _to_selection(field: Selection, value: Any) -> str:
    """Значение варианта или его подпись, без учёта регистра."""
    word = text(value).lower()
    for option_value, label in field.options:
        if word in (option_value.lower(), label.lower()):
            return option_value
    values = [option_value for option_value, _ in field.options]
    allowed = ", ".join(values[:10]) + ("…" if len(values) > 10 else "")
    raise ValueError(f"нет такого варианта; допустимо: {allowed}")
