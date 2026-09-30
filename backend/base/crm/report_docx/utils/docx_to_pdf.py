"""
Простая конвертация DOCX → PDF без LibreOffice: python-docx + fpdf2.

Запасной путь для контейнера, где LibreOffice нет (см. DocxReportEngine).
Переносит текст абзацев (жирный/курсив/подчёркивание, размер, выравнивание,
отступы), нумерацию списков, таблицы сеткой (объединения, границы, ширины
колонок), картинки в абзацах и ячейках, плавающие картинки (печать, подписи)
в точке их привязки и разрывы страниц. Точной вёрстки Word не даёт: нет
колонтитулов, цветов и фигур — для отчётов «текст + таблицы» этого хватает,
для полиграфии ставьте LibreOffice.

Нужен TTF-шрифт с кириллицей: DejaVu Sans (Linux — пакет fonts-dejavu-core,
он есть в docker/backend.Dockerfile) или Arial (Windows/macOS).
"""

import io
import logging
import os
import re
from typing import Any, Iterator

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsmap, qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from fpdf import FPDF
from fpdf.enums import CellBordersLayout, TextEmphasis
from fpdf.fonts import FontFace
from lxml import etree

log = logging.getLogger(__name__)

EMU_PER_PT = 12700
TWIPS_PER_PT = 20
DEFAULT_FONT_SIZE_PT = 11.0
LINE_HEIGHT = 1.15
# A4 в пунктах
DEFAULT_PAGE_SIZE_PT = (595.3, 841.9)
DEFAULT_MARGIN_PT = 56.7  # 2 см

# Семейства-кандидаты: (regular, bold, italic, bold-italic). Первое найденное.
_FONT_CANDIDATES = [
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf",
    ),
    (
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "C:/Windows/Fonts/ariali.ttf",
        "C:/Windows/Fonts/arialbi.ttf",
    ),
    (
        "/Library/Fonts/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf",
        "/Library/Fonts/Arial Italic.ttf",
        "/Library/Fonts/Arial Bold Italic.ttf",
    ),
    (
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf",
    ),
]

_ALIGN = {
    WD_ALIGN_PARAGRAPH.CENTER: "C",
    WD_ALIGN_PARAGRAPH.RIGHT: "R",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "J",
    WD_ALIGN_PARAGRAPH.DISTRIBUTE: "J",
}

_BORDER_SIDES = {
    "top": CellBordersLayout.TOP,
    "bottom": CellBordersLayout.BOTTOM,
    "left": CellBordersLayout.LEFT,
    "right": CellBordersLayout.RIGHT,
}

# Имена сторон в tblBorders/tcBorders → наши (start/end = left/right)
_BORDER_ALIASES = {
    "top": "top",
    "bottom": "bottom",
    "left": "left",
    "start": "left",
    "right": "right",
    "end": "right",
    "insideH": "insideH",
    "insideV": "insideV",
}

# Отступы ячейки: (верх, право, низ, лево)
CELL_PADDING = (2, 4, 2, 4)


def docx_to_pdf(docx_bytes: bytes, title: str | None = None) -> bytes:
    """Конвертирует DOCX в PDF. title — метаданные документа (имя файла
    при сохранении из просмотрщика). RuntimeError — не найден TTF-шрифт."""
    return _Converter(Document(io.BytesIO(docx_bytes)), title).run()


def _border_visible(border) -> bool:
    """Линия границы видна: val не none/nil, толщина не 0 и цвет не белый
    (белой сеткой Word-формы прячут границы у части ячеек)."""
    if border.get(qn("w:val")) in (None, "none", "nil"):
        return False
    if border.get(qn("w:sz")) == "0":
        return False
    return (border.get(qn("w:color")) or "").upper() != "FFFFFF"


def _style_chain(style) -> Iterator[Any]:
    while style is not None:
        yield style
        style = style.base_style


def _xpath(element, path: str) -> list:
    """XPath с пространствами имён Word. Нужен для элементов numbering.xml
    (w:lvl и т.п.), у которых нет обёртки python-docx с готовым nsmap."""
    return etree._Element.xpath(element, path, namespaces=nsmap)


def _format_number(value: int, fmt: str) -> str:
    """Номер пункта списка в формате numFmt (decimal, lowerLetter, upperRoman…)."""
    if fmt in ("lowerLetter", "upperLetter") and 1 <= value <= 26:
        letter = chr(ord("a") + value - 1)
        return letter.upper() if fmt == "upperLetter" else letter
    if fmt in ("lowerRoman", "upperRoman") and 1 <= value < 4000:
        roman = ""
        for number, symbol in (
            (1000, "m"),
            (900, "cm"),
            (500, "d"),
            (400, "cd"),
            (100, "c"),
            (90, "xc"),
            (50, "l"),
            (40, "xl"),
            (10, "x"),
            (9, "ix"),
            (5, "v"),
            (4, "iv"),
            (1, "i"),
        ):
            while value >= number:
                roman += symbol
                value -= number
        return roman.upper() if fmt == "upperRoman" else roman
    return str(value)


class _Converter:
    def __init__(self, doc, title: str | None = None):
        self.doc = doc
        self.pdf = self._make_pdf()
        if title:
            self.pdf.set_title(title)
        self.font_styles = self._register_fonts()
        self.default_size, self.default_before, self.default_after = (
            self._doc_defaults()
        )
        # Счётчики нумерованных списков: abstractNumId → {уровень: номер}
        self.list_counters: dict[int, dict[int, int]] = {}
        self.seen_num_ids: set[int] = set()

    # ------------------------------------------------------------------
    # Настройка страницы и шрифтов
    # ------------------------------------------------------------------

    def _make_pdf(self) -> FPDF:
        section = self.doc.sections[0] if self.doc.sections else None
        width, height = DEFAULT_PAGE_SIZE_PT
        left = right = top = bottom = DEFAULT_MARGIN_PT
        if section is not None:
            if section.page_width and section.page_height:
                width = section.page_width.pt
                height = section.page_height.pt
            left = section.left_margin.pt if section.left_margin else left
            right = section.right_margin.pt if section.right_margin else right
            top = section.top_margin.pt if section.top_margin else top
            bottom = (
                section.bottom_margin.pt if section.bottom_margin else bottom
            )
        pdf = FPDF(unit="pt", format=(width, height))
        pdf.set_margins(left, top, right)
        pdf.set_auto_page_break(True, margin=bottom)
        self.bottom_margin = bottom
        pdf.add_page()
        return pdf

    def _register_fonts(self) -> set[str]:
        """Регистрирует первое найденное семейство. Возвращает доступные
        начертания ("", "B", "I", "BI")."""
        for paths in _FONT_CANDIDATES:
            if not os.path.exists(paths[0]):
                continue
            styles: set[str] = set()
            for style, path in zip(("", "B", "I", "BI"), paths):
                if os.path.exists(path):
                    self.pdf.add_font("Doc", style, path)
                    styles.add(style)
            return styles
        raise RuntimeError(
            "PDF: не найден TTF-шрифт с кириллицей (DejaVu Sans / Arial). "
            "Linux: apt install fonts-dejavu-core"
        )

    def _doc_defaults(self) -> tuple[float, float, float]:
        """(размер шрифта, интервал до, интервал после) из docDefaults."""
        styles = self.doc.styles.element
        size = DEFAULT_FONT_SIZE_PT
        sz = styles.xpath("w:docDefaults/w:rPrDefault/w:rPr/w:sz")
        if sz and sz[0].get(qn("w:val")):
            size = int(sz[0].get(qn("w:val"))) / 2
        before = after = 0.0
        spacing = styles.xpath("w:docDefaults/w:pPrDefault/w:pPr/w:spacing")
        if spacing:
            before = int(spacing[0].get(qn("w:before")) or 0) / TWIPS_PER_PT
            after = int(spacing[0].get(qn("w:after")) or 0) / TWIPS_PER_PT
        return size, before, after

    def _set_font(self, run, paragraph) -> float:
        """Ставит шрифт по свойствам run (с учётом стилей). Возвращает размер."""
        size = self._run_prop(run, paragraph, "size")
        size = size.pt if size else self.default_size
        bold = self._run_prop(run, paragraph, "bold")
        italic = self._run_prop(run, paragraph, "italic")
        style = ("B" if bold else "") + ("I" if italic else "")
        if style not in self.font_styles:
            style = "B" if bold and "B" in self.font_styles else ""
        if self._run_prop(run, paragraph, "underline"):
            style += "U"
        self.pdf.set_font("Doc", style, size)
        return size

    @staticmethod
    def _run_prop(run, paragraph, name: str):
        """Свойство шрифта: сам run → стиль run → стиль абзаца (по цепочке)."""
        value = getattr(run.font, name)
        if value is not None:
            return value
        chain = list(_style_chain(run.style)) + list(
            _style_chain(paragraph.style)
        )
        for style in chain:
            value = getattr(style.font, name)
            if value is not None:
                return value
        return None

    @staticmethod
    def _para_prop(paragraph, name: str):
        """Свойство абзаца: сам абзац → стиль (по цепочке)."""
        value = getattr(paragraph.paragraph_format, name)
        if value is not None:
            return value
        for style in _style_chain(paragraph.style):
            value = getattr(style.paragraph_format, name)
            if value is not None:
                return value
        return None

    # ------------------------------------------------------------------
    # Обход документа
    # ------------------------------------------------------------------

    def run(self) -> bytes:
        for block in self.doc.iter_inner_content():
            if isinstance(block, Paragraph):
                self._paragraph(block)
            elif isinstance(block, Table):
                self._table(block)
        return bytes(self.pdf.output())

    @staticmethod
    def _runs(paragraph) -> list:
        """Run-ы абзаца, включая вложенные в гиперссылки."""
        runs = []
        for item in paragraph.iter_inner_content():
            runs.extend(item.runs if hasattr(item, "runs") else [item])
        return runs

    def _paragraph(self, paragraph: Paragraph) -> None:
        pdf = self.pdf
        runs = self._runs(paragraph)
        has_page_break = self._para_prop(
            paragraph, "page_break_before"
        ) or any(run._r.xpath('./w:br[@w:type="page"]') for run in runs)
        if has_page_break and pdf.y > pdf.t_margin:
            pdf.add_page()

        align = _ALIGN.get(self._para_prop(paragraph, "alignment"), "L")
        before = self._para_prop(paragraph, "space_before")
        after = self._para_prop(paragraph, "space_after")
        before = before.pt if before is not None else self.default_before
        after = after.pt if after is not None else self.default_after
        indent = self._para_prop(paragraph, "left_indent")
        first_line = self._para_prop(paragraph, "first_line_indent")
        spacing = self._para_prop(paragraph, "line_spacing")
        prefix, list_indent = self._list_prefix(paragraph)

        text_runs = [run for run in runs if run.text]
        images = [img for run in runs for img in self._images(run)]
        inline = [img for img in images if not img["anchor"]]
        # Плавающие картинки (печать, подписи без обтекания): рисуем в их
        # смещении от верха абзаца, поток текста не двигаем — как Word
        top = pdf.y + before
        for image in images:
            if image["anchor"]:
                self._place_anchored(image, top)
        if not text_runs:
            # Пустой абзац — пустая строка (так в Word делают отступы)
            size = (
                self._set_font(runs[0], paragraph)
                if runs
                else self.default_size
            )
            if not inline:
                pdf.ln(before + size * LINE_HEIGHT + after)
        else:
            size = self._set_font(text_runs[0], paragraph)
            if isinstance(spacing, (int, float)):
                line_height = float(spacing)
            elif spacing is not None:
                line_height = spacing.pt / size
            else:
                line_height = LINE_HEIGHT
            with pdf.text_columns(
                ncols=1,
                l_margin=pdf.l_margin + (indent.pt if indent else list_indent),
            ) as cols:
                with cols.paragraph(
                    text_align=align,
                    line_height=line_height,
                    top_margin=before,
                    bottom_margin=after,
                    first_line_indent=first_line.pt if first_line else 0,
                ) as par:
                    if prefix:
                        par.write(prefix)
                    for run in text_runs:
                        self._set_font(run, paragraph)
                        par.write(run.text.replace("\t", "    "))

        # Встроенные картинки — блоком после текста абзаца
        for image in inline:
            width, height = image["width"], image["height"]
            x = pdf.l_margin
            if align == "C":
                x = pdf.l_margin + (pdf.epw - width) / 2
            elif align == "R":
                x = pdf.w - pdf.r_margin - width
            try:
                pdf.image(io.BytesIO(image["blob"]), x=x, w=width, h=height)
            except Exception as e:  # EMF/WMF и прочее, что не открывает Pillow
                log.warning("PDF: картинка пропущена: %s", e)

    def _place_anchored(self, image: dict, top: float) -> None:
        """Плавающая картинка в точке её привязки: смещение по горизонтали от
        колонки/поля/страницы (или выравнивание), по вертикали — от верха
        абзаца, поля или страницы. Поток текста не двигает."""
        pdf = self.pdf
        base_x = 0.0 if image["rel_h"] == "page" else pdf.l_margin
        base_w = pdf.w if image["rel_h"] == "page" else pdf.epw
        if image["align_h"] == "center":
            x = base_x + (base_w - image["width"]) / 2
        elif image["align_h"] == "right":
            x = base_x + base_w - image["width"]
        else:
            x = base_x + image["x"]
        if image["rel_v"] == "page":
            base_y = 0.0
        elif image["rel_v"] == "margin":
            base_y = pdf.t_margin
        else:  # paragraph / line
            base_y = top
        y = base_y + image["y"]
        # Картинка может выйти за низ страницы — это её место, не перенос
        pdf.set_auto_page_break(False)
        try:
            pdf.image(
                io.BytesIO(image["blob"]),
                x=x,
                y=y,
                w=image["width"],
                h=image["height"],
            )
        except Exception as e:
            log.warning("PDF: плавающая картинка пропущена: %s", e)
        finally:
            pdf.set_auto_page_break(True, margin=self.bottom_margin)

    def _list_prefix(self, paragraph: Paragraph) -> tuple[str, float]:
        """Нумерация/маркер списка («1.2. », «• ») и отступ уровня в pt.

        Word хранит номера не в тексте, а в numbering.xml: считаем их сами
        по abstractNum (уровни, формат, шаблон «%1.%2.»). Продолжение
        нумерации между списками одного abstractNum — как в Word; перезапуск
        через startOverride учитываем при первой встрече numId.
        """
        # numPr самого абзаца, иначе — стиля (нумерованные заголовки)
        num_pr = None
        for p_pr in [paragraph._p.pPr] + [
            style.element.pPr for style in _style_chain(paragraph.style)
        ]:
            if p_pr is not None and p_pr.numPr is not None:
                num_pr = p_pr.numPr
                break
        if num_pr is None or num_pr.numId is None or not num_pr.numId.val:
            return "", 0.0
        try:
            numbering = self.doc.part.numbering_part.element
            num = numbering.num_having_numId(num_pr.numId.val)
        except (NotImplementedError, KeyError):
            return "", 0.0
        ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
        abstract_id = num.abstractNumId.val
        lvls = _xpath(
            numbering,
            f'w:abstractNum[@w:abstractNumId="{abstract_id}"]'
            f'/w:lvl[@w:ilvl="{ilvl}"]',
        )
        if not lvls:
            return "", 0.0
        lvl = lvls[0]
        fmt = (_xpath(lvl, "w:numFmt/@w:val") or ["decimal"])[0]
        template = (_xpath(lvl, "w:lvlText/@w:val") or ["%1."])[0]
        start = int((_xpath(lvl, "w:start/@w:val") or [1])[0])
        indent = int((_xpath(lvl, "w:pPr/w:ind/@w:left") or [0])[0])

        counters = self.list_counters.setdefault(abstract_id, {})
        if num_pr.numId.val not in self.seen_num_ids:
            self.seen_num_ids.add(num_pr.numId.val)
            if _xpath(num, "w:lvlOverride/w:startOverride"):
                counters.clear()
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for level in [level for level in counters if level > ilvl]:
            del counters[level]

        if fmt == "bullet":
            return "• ", indent / TWIPS_PER_PT
        text = re.sub(
            r"%(\d)",
            lambda m: _format_number(
                counters.get(int(m.group(1)) - 1, 1), fmt
            ),
            template,
        )
        return text + " ", indent / TWIPS_PER_PT

    def _images(self, run) -> list[dict]:
        """Картинки run-а: blob, width/height (pt), anchor (плавающая) и её
        привязка: x/y (pt), rel_h/rel_v (column|margin|page|paragraph|line),
        align_h (left|center|right)."""
        result = []
        for drawing in _xpath(run._r, "./w:drawing"):
            rids = _xpath(drawing, ".//a:blip/@r:embed")
            extents = _xpath(drawing, ".//wp:extent")
            if not rids or not extents:
                continue
            part = run.part.related_parts.get(rids[0])
            if part is None:
                continue
            image = {
                "blob": part.blob,
                "width": int(extents[0].get("cx")) / EMU_PER_PT,
                "height": int(extents[0].get("cy")) / EMU_PER_PT,
                "anchor": False,
                "x": 0.0,
                "y": 0.0,
                "rel_h": "column",
                "rel_v": "paragraph",
                "align_h": None,
            }
            anchors = _xpath(drawing, "wp:anchor")
            if anchors:
                image["anchor"] = True
                for axis, key, rel_key in (
                    ("wp:positionH", "x", "rel_h"),
                    ("wp:positionV", "y", "rel_v"),
                ):
                    positions = _xpath(anchors[0], axis)
                    if not positions:
                        continue
                    position = positions[0]
                    image[rel_key] = (
                        position.get("relativeFrom") or image[rel_key]
                    )
                    offsets = _xpath(position, "wp:posOffset")
                    aligns = _xpath(position, "wp:align")
                    if offsets and offsets[0].text:
                        image[key] = int(offsets[0].text) / EMU_PER_PT
                    if key == "x" and aligns:
                        image["align_h"] = aligns[0].text
            result.append(image)
        return result

    # ------------------------------------------------------------------
    # Таблицы
    # ------------------------------------------------------------------

    def _table(self, table: Table) -> None:
        pdf = self.pdf
        ncols = len(table.columns)
        if not ncols or not table.rows:
            return
        col_widths = [
            col.width.pt if col.width else None for col in table.columns
        ]
        widths = None
        table_width = None
        if all(col_widths):
            widths = tuple(col_widths)
            table_width = min(sum(col_widths), pdf.epw)

        size = self._table_font_size(table)
        pdf.set_font("Doc", "", size)
        align = {1: "CENTER", 2: "RIGHT"}.get(
            int(table.alignment or 0), "LEFT"
        )
        # Границы считаем по ячейкам: стороны из tblBorders таблицы/стиля
        # (снаружи — top/bottom/left/right, внутри — insideH/insideV), поверх
        # них — tcBorders самой ячейки. Белые и нулевые линии — не границы:
        # так в формах Word прячут сетку у части ячеек.
        table_sides = self._table_border_sides(table)
        # Ширины колонок в pt — чтобы картинку в ячейке рисовать в её размере
        total = table_width or pdf.epw
        shares = widths or tuple([1] * ncols)
        col_pts = [total * share / sum(shares) for share in shares]
        # Одна сетка ячеек на таблицу: объединённые python-docx отдаёт тем же
        # объектом (по горизонтали — подряд, по вертикали — из строки выше)
        grid = table._cells
        nrows = len(table.rows)

        with pdf.table(
            first_row_as_headings=False,
            col_widths=widths,
            width=table_width,
            align=align,
            line_height=size * LINE_HEIGHT,
            padding=CELL_PADDING,
            borders_layout="NONE",
        ) as pdf_table:
            for r, row in enumerate(table.rows):
                pdf_row = pdf_table.row()
                filled = 0
                row_cells = grid[r * ncols : (r + 1) * ncols]
                for cell, colspan in self._row_cells(row_cells):
                    if filled + colspan > ncols:
                        break
                    # Продолжение вертикального объединения: python-docx
                    # отдаёт ячейку верхней строки — текст не повторяем
                    continuation = cell._tc.getparent() is not row._tr
                    continued_below = r + 1 < nrows and any(
                        below is cell
                        for below in grid[(r + 1) * ncols : (r + 2) * ncols]
                    )
                    border = self._cell_border_layout(
                        cell,
                        table_sides,
                        edges=(
                            r == 0,
                            r == nrows - 1,
                            filled == 0,
                            filled + colspan >= ncols,
                        ),
                        merged=(continuation, continued_below),
                    )
                    text = "" if continuation else cell.text
                    images = [] if continuation else self._cell_images(cell)
                    if images and not text.strip():
                        blob, img_w = images[0]["blob"], images[0]["width"]
                        col_w = sum(col_pts[filled : filled + colspan])
                        top, right, bottom, left = CELL_PADDING
                        pdf_row.cell(
                            img=io.BytesIO(blob),
                            img_fill_width=True,
                            colspan=colspan,
                            border=border,
                            # Картинку в её размере: добираем правым отступом
                            padding=(
                                top,
                                max(right, col_w - left - img_w),
                                bottom,
                                left,
                            ),
                        )
                    else:
                        pdf_row.cell(
                            text,
                            colspan=colspan,
                            align=self._cell_align(cell),
                            style=self._cell_face(cell, size),
                            border=border,
                        )
                    filled += colspan
                for _ in range(ncols - filled):
                    pdf_row.cell("", border=CellBordersLayout.NONE)

    def _cell_images(self, cell) -> list[dict]:
        return [
            image
            for paragraph in cell.paragraphs
            for run in self._runs(paragraph)
            for image in self._images(run)
        ]

    @staticmethod
    def _table_border_sides(table: Table) -> dict[str, bool]:
        """Стороны tblBorders таблицы (свои поверх стиля, по каждой стороне
        отдельно): top/bottom/left/right — внешние, insideH/insideV —
        внутренние линии. Не заданная нигде сторона — без линии."""
        sides: dict[str, bool] = {}
        sources = [table._tbl.tblPr] + [
            style.element for style in _style_chain(table.style)
        ]
        for element in sources:
            if element is None:
                continue
            for border in element.xpath(".//w:tblBorders/*"):
                side = _BORDER_ALIASES.get(border.tag.rsplit("}", 1)[-1])
                if side and side not in sides:
                    sides[side] = _border_visible(border)
        return sides

    @staticmethod
    def _cell_border_layout(
        cell,
        table_sides: dict[str, bool],
        edges: tuple[bool, bool, bool, bool],
        merged: tuple[bool, bool],
    ) -> CellBordersLayout:
        """Границы ячейки: линии таблицы по её положению (edges: первая/
        последняя строка, первая/последняя колонка), поверх — tcBorders.
        merged: (продолжение вертикального объединения, продолжается ниже) —
        внутри объединённой ячейки линий нет."""
        first_row, last_row, first_col, last_col = edges
        visible = {
            "top": table_sides.get("top" if first_row else "insideH", False),
            "bottom": table_sides.get(
                "bottom" if last_row else "insideH", False
            ),
            "left": table_sides.get("left" if first_col else "insideV", False),
            "right": table_sides.get(
                "right" if last_col else "insideV", False
            ),
        }
        tc_pr = cell._tc.tcPr
        if tc_pr is not None:
            for border in tc_pr.xpath("w:tcBorders/*"):
                side = _BORDER_ALIASES.get(border.tag.rsplit("}", 1)[-1])
                if side in visible:
                    visible[side] = _border_visible(border)
        continuation, continued_below = merged
        if continuation:
            visible["top"] = False
        if continued_below:
            visible["bottom"] = False
        layout = CellBordersLayout.NONE
        for side, flag in _BORDER_SIDES.items():
            if visible[side]:
                layout |= flag
        return layout

    @staticmethod
    def _row_cells(cells) -> list[tuple[Any, int]]:
        """Ячейки строки с colspan: объединённые по горизонтали python-docx
        отдаёт повтором одного объекта."""
        result: list[tuple[Any, int]] = []
        for cell in cells:
            if result and result[-1][0] is cell:
                result[-1] = (cell, result[-1][1] + 1)
            else:
                result.append((cell, 1))
        return result

    def _table_font_size(self, table: Table) -> float:
        """Размер шрифта таблицы — по первому непустому run-у."""
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in self._runs(paragraph):
                        if run.text.strip():
                            size = self._run_prop(run, paragraph, "size")
                            return size.pt if size else self.default_size
        return self.default_size

    def _cell_align(self, cell) -> str:
        paragraphs = cell.paragraphs
        if not paragraphs:
            return "L"
        return _ALIGN.get(self._para_prop(paragraphs[0], "alignment"), "L")

    def _cell_face(self, cell, table_size: float) -> FontFace | None:
        """Жирный — если весь текст ячейки жирный; размер — если отличается."""
        runs = [
            run
            for paragraph in cell.paragraphs
            for run in self._runs(paragraph)
            if run.text.strip()
        ]
        if not runs:
            return None
        paragraph = cell.paragraphs[0]
        bold = all(self._run_prop(run, paragraph, "bold") for run in runs)
        size = self._run_prop(runs[0], paragraph, "size")
        size_pt = size.pt if size else self.default_size
        if not bold and size_pt == table_size:
            return None
        return FontFace(
            emphasis=(
                TextEmphasis.B if bold and "B" in self.font_styles else None
            ),
            size_pt=size_pt if size_pt != table_size else None,
        )
